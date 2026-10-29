"""指定テンプレート SupportingData_format_V2.pptx に、図 3 点（列）と文章を入れて SupportingData_filled_<する|しない>.pptx を作る。
usage: python3 -I docs/jsmp133/supporting_data/build_pptx.py [する|しない]   （python-pptx が要る。venv には入っていない）
図は build_figures.py で作る（fig_method.png / fig_result.png / fig_apps.png、各 3.27×4.30 インチ）。
注意: 図形を複製すると id が重複して Keynote で開けない（add_textbox で作る）。Keynote の書き出しは abstract_plan.md の「Supporting Data」参照。
匿名審査のため、著者名・所属・施設名・謝辞は入れない。
"""
import sys, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt, Inches

D = Path(__file__).resolve().parent
entry = sys.argv[1] if len(sys.argv) > 1 else 'する'
assert entry in ('する', 'しない')

prs = Presentation(str(D / 'SupportingData_format_V2.pptx'))
sl = prs.slides[0]
by_id = {sh.shape_id: sh for sh in sl.shapes}


def drop(sh):
    sh._element.getparent().remove(sh._element)


def set_text(sh, text):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# テンプレートの見本（IMRT の例の図・赤い注記・矢印）を外す
for i in (31, 35, 6, 7, 8, 9):
    drop(by_id[i])

set_text(by_id[21], 'AI対話型モンテカルロシミュレーションシステムの開発')
set_text(by_id[2], f'賞にエントリー：{entry}')

# 図: 3 列
W, H = Inches(3.27), Inches(4.30)
LEFT, TOP = 378120, 1175000
GAP = (9192288 - 3 * W) // 2
cols = [('fig_method.png', 'fig_method'), ('fig_result.png', 'fig_result'), ('fig_apps.png', 'fig_apps')]
xs = [LEFT + i * (W + GAP) for i in range(3)]
for (f, _), x in zip(cols, xs):
    sl.shapes.add_picture(str(D / f), x, TOP, W, H)

# 図表の説明（テンプレートの「図表の説明」2 箱 + 複製 1 箱）
CAP_TOP = TOP + H + 15000
CAP_H = 190000
cap3 = sl.shapes.add_textbox(xs[2], CAP_TOP, W, CAP_H)
cap_shapes = [by_id[3], by_id[4], cap3]
caps = [
    '図1　計算の流れと、2 つの障壁への対応。',
    '図2　差は EGS5 値に対する割合（OCR は照射野外含む）',
    '図3　いずれも対話で実施。X 線画像は共通の濃淡。',
]
for sh, x, text in zip(cap_shapes, xs, caps):
    sh.left, sh.top, sh.width, sh.height = x, CAP_TOP, W, CAP_H
    tf = sh.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr')
    bp.set('wrap', 'square')
    for ch in list(bp):
        bp.remove(ch)
    p0 = tf.paragraphs[0]
    r = p0.runs[0] if p0.runs else p0.add_run()
    r.text = text
    r.font.size = Pt(8)

# イノベーション欄（帯と文章を下へ移す）
BAND_TOP = CAP_TOP + CAP_H + 45000
delta = BAND_TOP - by_id[17].top
by_id[17].top = Emu(BAND_TOP)
by_id[18].top = Emu(by_id[18].top + delta)
box = by_id[32]
box.top = Emu(BAND_TOP + 390380 + 20000)
para = box.text_frame.paragraphs[0]
para.line_spacing = 0.95
for r in list(para.runs):
    r._r.getparent().remove(r._r)
SEGS = [
    ('従来のモンテカルロ計算の 2 つの障壁を、', False),
    ('AI エージェントによる条件設定の支援（入力形式の理解）と、Google Colaboratory のクラウド GPU（計算資源）で下げた', True),
    ('点に新規性がある。使用者は 3D 表示で条件を確認・承認して、意図した計算を任意の形式で得られる。EGS5 との差は', False),
    ('平均 0.4% 未満、最大 1.3%', True),
    ('で、性質の異なる依頼（X 線画像、前立腺シード）も対話で実施できた。', False),
    ('手順書形式（Agent Skills）により特定の AI に依存せず、研究・教育の場への波及が期待される', True),
    ('。', False),
]
for text, bold in SEGS:
    r = para.add_run()
    r.text = text
    r.font.size = Pt(11.5)
    r.font.bold = bold
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', 'ja-JP')
    rPr.set('altLang', 'en-US')

out = D / f'SupportingData_filled_{entry}.pptx'
prs.save(str(out))
print('saved', out)
